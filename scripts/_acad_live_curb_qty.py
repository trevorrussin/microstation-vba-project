"""Live curb LF -> Pre-Final Cost Estimate Working (4.09 AD/BD/CD only).

Target workbook (in-place Length updates, never deleted):
  2026.09.28_HWP24QC1 - Pre-Final Cost Estimate_Working.xlsx

Writes Length (col D) on:
  4.09 AD = Straight
  4.09 BD = Depressed
  4.09 CD = Corner
matched by Corner ID (col C).

Does NOT write HW-200 or C.E. (Main) — those VLOOKUP/ref 4.09.

Corner geography is project-wide (all locations), not Loc01-only.
Regions live in Data/corner-cad-regions.json (add centers as you classify
more corners). Optional: XData app "HWP24QC1" key "CORNER_ID" on a line
overrides spatial assignment.

Usage:
  python scripts/_acad_live_curb_qty.py
  python scripts/_acad_live_curb_qty.py --watch
  python scripts/_acad_live_curb_qty.py --dry-run
"""
from __future__ import annotations

import argparse
import json
import math
import os
import sys
import time
from pathlib import Path

import pythoncom
from win32com.client import Dispatch, GetObject

ROOT = Path(r"c:\repos\microstation-vba-project")
REGIONS_PATH = ROOT / "Data" / "corner-cad-regions.json"
DASH_CSV = ROOT / "outputs" / "curb-live-qty.csv"
WORKING_XLSX = ROOT / "2026.09.28_HWP24QC1 - Pre-Final Cost Estimate_Working.xlsx"

LAYERS = {
    "Straight Steel Faced Concrete Curb (18'' Deep)": "4.09 AD",
    "Depressed Steel Faced Concrete Curb (18'' Deep)": "4.09 BD",
    "Corner Steel Faced Concrete Curb (18'' Deep)": "4.09 CD",
}
XDATA_APP = "HWP24QC1"


def load_regions():
    if not REGIONS_PATH.exists():
        return {"corners": {}}
    return json.loads(REGIONS_PATH.read_text(encoding="utf-8"))


def excel_corner_index(path):
    """All Corner IDs from 4.09 AD (Location, Corner ID, row)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ws = wb["4.09 AD"]
    rows = {}
    for r, row in enumerate(
        ws.iter_rows(min_row=14, max_row=200, min_col=2, max_col=3, values_only=True),
        14,
    ):
        loc, cid = row[0], row[1]
        if cid is None:
            continue
        try:
            cid_i = int(cid)
        except (TypeError, ValueError):
            continue
        loc_i = int(loc) if loc is not None and str(loc).strip() != "" else None
        rows[cid_i] = {"location_no": loc_i, "row": r}
    wb.close()
    return rows


def read_xdata_corner(e):
    """Return int Corner ID from XData if present."""
    try:
        # AutoCAD COM: GetXData(appName) -> (types, values)
        t, v = e.GetXData(XDATA_APP)
        if not t:
            return None
        # pairs: 1001 app, 1000 string / 1071 int32 etc.
        for typ, val in zip(t, v):
            if typ in (1070, 1071) and val is not None:
                return int(val)
            if typ == 1000 and val is not None:
                s = str(val).strip()
                if s.isdigit():
                    return int(s)
    except Exception:
        return None
    return None


def nearest_corner(cx, cy, regions):
    best = None
    for cid_s, meta in regions.get("corners", {}).items():
        c = meta.get("center")
        if not c or len(c) < 2:
            continue
        r = float(meta.get("radius_ft", 35.0))
        d = math.hypot(cx - float(c[0]), cy - float(c[1]))
        if d <= r and (best is None or d < best[0]):
            best = (d, int(cid_s))
    return best[1] if best else None


def collect_lf_by_corner(regions):
    """{corner_id: {sheet: lf}} from AutoCAD steel lines."""
    pythoncom.CoInitialize()
    doc = GetObject(None, "AutoCAD.Application").ActiveDocument
    ms = doc.ModelSpace
    out = {}
    n_ok = n_skip = 0
    for i in range(ms.Count):
        try:
            e = ms.Item(i)
            if e.ObjectName != "AcDbLine" or e.Layer not in LAYERS:
                continue
            sheet = LAYERS[e.Layer]
            mn, mx = e.GetBoundingBox()
            cx = (float(mn[0]) + float(mx[0])) / 2
            cy = (float(mn[1]) + float(mx[1])) / 2
            cid = read_xdata_corner(e) or nearest_corner(cx, cy, regions)
            if cid is None:
                n_skip += 1
                continue
            slot = out.setdefault(cid, {s: 0.0 for s in LAYERS.values()})
            slot[sheet] += float(e.Length)
            n_ok += 1
        except Exception:
            continue
    return out, n_ok, n_skip, doc.Name


def write_dashboard(by_corner, excel_idx):
    DASH_CSV.parent.mkdir(parents=True, exist_ok=True)
    lines = ["corner_id,location_no,sheet,lf"]
    for cid, sheets in sorted(by_corner.items()):
        loc = excel_idx.get(cid, {}).get("location_no")
        for sh, lf in sheets.items():
            lines.append(f"{cid},{loc},{sh},{lf:.4f}")
    DASH_CSV.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_openpyxl(path, by_corner, excel_idx, dry_run=False):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    updates = []
    for cid, sheets in by_corner.items():
        info = excel_idx.get(cid)
        if not info:
            updates.append(("?", cid, None, None, "CORNER_NOT_IN_4.09_AD"))
            continue
        row = info["row"]
        for sheet_name, lf in sheets.items():
            if sheet_name not in ("4.09 AD", "4.09 BD", "4.09 CD"):
                raise RuntimeError(f"refusing sheet {sheet_name}")
            ws = wb[sheet_name]
            # confirm Corner ID on this row
            cell_cid = ws.cell(row, 3).value
            try:
                if int(cell_cid) != int(cid):
                    updates.append((sheet_name, cid, row, lf, "ROW_MISMATCH"))
                    continue
            except (TypeError, ValueError):
                updates.append((sheet_name, cid, row, lf, "BAD_CID"))
                continue
            old = ws.cell(row, 4).value
            updates.append((sheet_name, cid, row, lf, old))
            if not dry_run:
                ws.cell(row, 4).value = round(lf, 2)
    if not dry_run:
        # save in place — never delete the Working file
        wb.save(path)
    wb.close()
    return updates


def write_excel_com(path, by_corner, excel_idx, dry_run=False):
    pythoncom.CoInitialize()
    try:
        xl = Dispatch("Excel.Application")
    except Exception:
        return None
    wb = None
    base = os.path.basename(path).lower()
    try:
        for i in range(1, xl.Workbooks.Count + 1):
            w = xl.Workbooks(i)
            if w.Name.lower() == base or w.FullName.lower() == path.lower():
                wb = w
                break
    except Exception:
        return None
    if wb is None:
        return None
    updates = []
    for cid, sheets in by_corner.items():
        info = excel_idx.get(cid)
        if not info:
            updates.append(("?", cid, None, None, "CORNER_NOT_IN_4.09_AD"))
            continue
        row = info["row"]
        for sheet_name, lf in sheets.items():
            ws = wb.Worksheets(sheet_name)
            old = ws.Cells(row, 4).Value
            updates.append((sheet_name, cid, row, lf, old))
            if not dry_run:
                ws.Cells(row, 4).Value = round(lf, 2)
    if not dry_run:
        try:
            xl.CalculateFull()
        except Exception:
            pass
    return updates


def run_once(dry_run=False):
    path = WORKING_XLSX
    if not path.exists():
        print("MISSING Working estimate:", path)
        return 1
    print("TARGET workbook (in-place):", path)
    print("  (HW-200 / C.E. Main not written — VLOOKUP from 4.09)")

    regions = load_regions()
    excel_idx = excel_corner_index(str(path))
    print(f"4.09 AD Corner IDs in Working estimate: {len(excel_idx)}")
    print(f"CAD regions registered: {len(regions.get('corners', {}))}")

    by_corner, n_ok, n_skip, dwg = collect_lf_by_corner(regions)
    print(f"DWG: {dwg}  assigned_segs={n_ok}  unassigned_steel={n_skip}")
    write_dashboard(by_corner, excel_idx)

    if not by_corner:
        print("No LF assigned — add corner centers to Data/corner-cad-regions.json")
        print("or stamp XData CORNER_ID on steel lines.")
        return 0

    print("--- LF by Corner ID ---")
    for cid, sheets in sorted(by_corner.items()):
        loc = excel_idx.get(cid, {}).get("location_no")
        print(
            f"  Loc {loc} corner {cid}: "
            + "  ".join(f"{s}={v:.2f}" for s, v in sheets.items())
        )

    updates = write_excel_com(str(path), by_corner, excel_idx, dry_run=dry_run)
    mode = "Excel COM (open Working book)"
    if updates is None:
        updates = write_openpyxl(str(path), by_corner, excel_idx, dry_run=dry_run)
        mode = "openpyxl save-in-place"
    print(f"--- writes via {mode} dry_run={dry_run} ---")
    for sheet_name, cid, row, lf, old in updates:
        print(f"  {sheet_name} corner={cid} row={row} {old} -> {lf}")
    print("dashboard", DASH_CSV)
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--watch", action="store_true")
    ap.add_argument("--interval", type=float, default=1.0)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    if not args.watch:
        return run_once(dry_run=args.dry_run)
    print(f"WATCH every {args.interval}s -> {WORKING_XLSX.name}  Ctrl+C stop")
    while True:
        try:
            run_once(dry_run=args.dry_run)
        except Exception as ex:
            print("tick error:", ex)
        print("---")
        time.sleep(args.interval)


if __name__ == "__main__":
    sys.exit(main() or 0)
