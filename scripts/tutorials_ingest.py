"""Extract text from Standards/Tutorials (PDF, PPTX, XLSX, MP4) into Standards/derived/tutorials/.

  python scripts/tutorials_ingest.py            # everything not yet done (resumable)
  python scripts/tutorials_ingest.py --only "Cutting Profiles"
  python scripts/tutorials_ingest.py --docs-only

Videos: local GPU speech-to-text (faster-whisper) with timestamps + on-screen text (OCR of scene-change frames:
dialog titles, tool names, key-ins) merged into one timeline:  [mm:ss] SAY: ...  /  [mm:ss] SCREEN: ...
Every output file starts with its source path so anything distilled from it can cite file + timestamp.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time

try:  # corporate TLS interception (model download)
    import truststore
    truststore.inject_into_ssl()
except Exception:
    pass

def _add_cuda_dll_dirs():  # pip nvidia-* wheels (cublas / cudnn) for ctranslate2; torch here is the CPU build
    import glob
    import site
    for base in site.getsitepackages():
        for d in glob.glob(os.path.join(base, "nvidia", "*", "bin")):
            try:
                os.add_dll_directory(d)
            except Exception:
                pass
            os.environ["PATH"] = d + os.pathsep + os.environ.get("PATH", "")


_add_cuda_dll_dirs()
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "Standards", "Tutorials")
OUT = os.path.join(ROOT, "Standards", "derived", "tutorials")
MODEL = os.environ.get("TUTORIAL_WHISPER_MODEL", "large-v3")
PROMPT = ("MicroStation, OpenRoads Designer, ORD, NYSDOT, alignment, profile, corridor, terrain, cross section, "
          "named boundary, civil cell, level, reference file, model view, key-in, line style, hatch, stationing, "
          "plan sheet, geometry, template, superelevation, ProjectWise.")


def slug(path: str) -> str:
    rel = os.path.relpath(path, SRC)
    return re.sub(r"[^A-Za-z0-9]+", "_", os.path.splitext(rel)[0]).strip("_")[:120]


def stamp(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 3600:d}:{sec % 3600 // 60:02d}:{sec % 60:02d}" if sec >= 3600 else f"{sec // 60:02d}:{sec % 60:02d}"


def write(out_dir: str, name: str, header: str, body: str) -> None:
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, name), "w", encoding="utf-8") as f:
        f.write(header + "\n\n" + body)


_ocr = None


def ocr_image(img) -> list[str]:
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    res, _ = _ocr(img)
    return [r[1] for r in (res or []) if len(r[1].strip()) > 2 and float(r[2]) > 0.6]


# ------------------------------------------------------------------ documents
def do_pdf(path, out):
    import fitz
    parts = []
    for p in fitz.open(path):
        txt = p.get_text().strip()
        if len(txt) < 40:  # scanned / image page
            pix = p.get_pixmap(dpi=150)
            import numpy as np
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, :3]
            txt = "\n".join(ocr_image(img[:, :, ::-1].copy()))
        parts.append(f"## page {p.number + 1}\n{txt}")
    write(out, "text.md", f"# SOURCE: {os.path.relpath(path, ROOT)}", "\n\n".join(parts))


def _raw_zip_members(path):
    """Tolerant zip reader for a damaged central directory: scan the local file headers (read-only; never touches the file)."""
    import struct
    import zlib
    data = open(path, "rb").read()
    members, i = {}, 0
    while True:
        i = data.find(b"PK", i)
        if i < 0:
            break
        try:
            _, _, _, comp, _, _, _, csz, _, nlen, elen = struct.unpack("<IHHHHHIIIHH", data[i:i + 30])
            name = data[i + 30:i + 30 + nlen].decode("utf-8", "replace")
            members[name] = (comp, data[i + 30 + nlen + elen:i + 30 + nlen + elen + (csz or 4_000_000)])
        except Exception:
            pass
        i += 4

    def read(name):
        comp, blob = members[name]
        return zlib.decompressobj(-15).decompress(blob) if comp == 8 else blob
    return members, read


def do_pptx_recovered(path, out):
    import re
    members, read = _raw_zip_members(path)
    slides = sorted((n for n in members if re.match(r"ppt/slides/slide\d+\.xml$", n)), key=lambda n: int(re.findall(r"\d+", n)[0]))
    parts = []
    for n in slides:
        num = re.findall(r"\d+", n)[0]
        xml = read(n).decode("utf-8", "replace")
        lines = [t.strip() for t in re.findall(r"<a:t>([^<]*)</a:t>", xml) if t.strip()]
        rels = f"ppt/slides/_rels/slide{num}.xml.rels"
        if rels in members:
            for tgt in re.findall(r'Target="\.\./media/([^"]+)"', read(rels).decode("utf-8", "replace")):
                try:
                    import cv2
                    import numpy as np
                    img = cv2.imdecode(np.frombuffer(read("ppt/media/" + tgt), np.uint8), cv2.IMREAD_COLOR)
                    if img is not None and min(img.shape[:2]) > 60:
                        txt = ocr_image(img)
                        if txt:
                            lines.append("[image text] " + " / ".join(txt))
                except Exception:
                    pass
        notes = f"ppt/notesSlides/notesSlide{num}.xml"
        if notes in members:
            nt = " ".join(t.strip() for t in re.findall(r"<a:t>([^<]*)</a:t>", read(notes).decode("utf-8", "replace")) if t.strip())
            if nt:
                lines.append("[notes] " + nt)
        parts.append(f"## slide {num}\n" + "\n".join(lines))
    write(out, "text.md", f"# SOURCE: {os.path.relpath(path, ROOT)}  (recovered from a damaged zip index)", "\n\n".join(parts))


def do_pptx(path, out):
    try:
        return _do_pptx(path, out)
    except OSError:
        return do_pptx_recovered(path, out)


def _do_pptx(path, out):
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:  # group
                    walk(sh.shapes)
                    continue
                if sh.has_text_frame and sh.text_frame.text.strip():
                    lines.append(sh.text_frame.text.strip())
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        lines.append(" | ".join(c.text.strip() for c in row.cells))
                if sh.shape_type == 13:  # picture -> OCR (screenshots carry the tips)
                    try:
                        import numpy as np
                        import cv2
                        arr = cv2.imdecode(np.frombuffer(sh.image.blob, np.uint8), cv2.IMREAD_COLOR)
                        if arr is not None and min(arr.shape[:2]) > 60:
                            t = ocr_image(arr)
                            if t:
                                lines.append("[image text] " + " / ".join(t))
                    except Exception:
                        pass
        walk(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("[notes] " + slide.notes_slide.notes_text_frame.text.strip())
        parts.append(f"## slide {i}\n" + "\n".join(lines))
    write(out, "text.md", f"# SOURCE: {os.path.relpath(path, ROOT)}", "\n\n".join(parts))


def do_xlsx(path, out):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        parts.append(f"## sheet {ws.title}\n" + "\n".join(rows))
    write(out, "text.md", f"# SOURCE: {os.path.relpath(path, ROOT)}", "\n\n".join(parts))


# ------------------------------------------------------------------ video
def screen_text(path, out, every=3.0, min_gap=6.0, thumb_every=30.0):
    """OCR scene-change frames. Returns [(t, [lines])] with lines already seen recently removed."""
    import cv2
    import numpy as np
    cap = cv2.VideoCapture(path)
    fps = cap.get(cv2.CAP_PROP_FPS) or 25
    n = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    step = max(int(fps * every), 1)
    last_small, last_ocr_t, last_thumb = None, -999.0, -999.0
    seen: dict[str, float] = {}
    res = []
    os.makedirs(os.path.join(out, "frames"), exist_ok=True)
    for fi in range(0, n, step):
        cap.set(cv2.CAP_PROP_POS_FRAMES, fi)
        ok, frame = cap.read()
        if not ok:
            break
        t = fi / fps
        small = cv2.cvtColor(cv2.resize(frame, (192, 108)), cv2.COLOR_BGR2GRAY)
        changed = last_small is None or float(np.mean(cv2.absdiff(small, last_small))) > 4.0
        last_small = small
        if t - last_thumb >= thumb_every:
            cv2.imwrite(os.path.join(out, "frames", f"{int(t):05d}.jpg"), cv2.resize(frame, (1280, 720)), [cv2.IMWRITE_JPEG_QUALITY, 70])
            last_thumb = t
        if not changed or t - last_ocr_t < min_gap:
            continue
        last_ocr_t = t
        lines = [ln.strip() for ln in ocr_image(frame)]
        fresh = [ln for ln in lines if t - seen.get(ln, -999) > 60]
        for ln in lines:
            seen[ln] = t
        if fresh:
            res.append((t, fresh))
    return res


_model = None


def transcribe(path):
    global _model
    if _model is None:
        from faster_whisper import WhisperModel
        _model = WhisperModel(MODEL, device="cuda", compute_type="float16")
    # condition_on_previous_text=False + a single temperature stop the retry/hallucination loops that stalled long,
    # noisy meeting recordings (2026-10-23); progress lines show it is moving.
    segs, info = _model.transcribe(path, language="en", vad_filter=True, beam_size=3, initial_prompt=PROMPT,
                                   condition_on_previous_text=False, temperature=0.0)
    out, last = [], -1
    for s in segs:
        if s.text.strip():
            out.append((s.start, s.text.strip()))
        if int(s.start // 120) != last:
            last = int(s.start // 120)
            print(f"      speech {stamp(s.start)} / {stamp(info.duration)}", flush=True)
    return out, info.duration


def do_video(path, out, ocr=True):
    t0 = time.time()
    said, dur = transcribe(path)
    silent = len(said) == 0  # screen-recording without narration: the steps ARE the frames -> sample densely
    scr = (screen_text(path, out, every=2.0, min_gap=3.0, thumb_every=4.0) if silent else screen_text(path, out)) if ocr else []
    items = [(t, "SAY", x) for t, x in said] + [(t, "SCREEN", " / ".join(ls)) for t, ls in scr]
    items.sort(key=lambda r: (r[0], r[1] != "SAY"))
    body = "\n".join(f"[{stamp(t)}] {kind}: {txt}" for t, kind, txt in items)
    write(out, "transcript.md", f"# SOURCE: {os.path.relpath(path, ROOT)}  (duration {stamp(dur)}; whisper {MODEL} + frame OCR)", body)
    with open(os.path.join(out, "done.json"), "w") as f:
        json.dump({"seconds": round(time.time() - t0), "say": len(said), "screen": len(scr), "silent": silent}, f)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="")
    ap.add_argument("--docs-only", action="store_true")
    ap.add_argument("--videos-only", action="store_true")
    ap.add_argument("--no-ocr", action="store_true")
    ap.add_argument("--redo", action="store_true", help="overwrite outputs that already exist")
    ap.add_argument("--reverse", action="store_true", help="work from the end of the list (second worker)")
    a = ap.parse_args()
    files = []
    for dp, _, fns in os.walk(SRC):
        for fn in sorted(fns):
            if fn.lower().endswith((".pdf", ".pptx", ".xlsx", ".mp4")):
                files.append(os.path.join(dp, fn))
    for path in sorted(files, reverse=a.reverse):
        if a.only and a.only.lower() not in path.lower():
            continue
        ext = os.path.splitext(path)[1].lower()
        if (a.docs_only and ext == ".mp4") or (a.videos_only and ext != ".mp4"):
            continue
        out = os.path.join(OUT, slug(path))
        marker = os.path.join(out, "done.json" if ext == ".mp4" else "text.md")
        if os.path.exists(marker) and not a.redo:
            continue
        if ext == ".mp4" and os.path.exists(os.path.join(out, "started.json")) and not a.redo:
            continue  # another worker has it
        if ext == ".mp4":
            os.makedirs(out, exist_ok=True)
            open(os.path.join(out, "started.json"), "w").write("{}")
        print(f"[{time.strftime('%H:%M:%S')}] {os.path.relpath(path, SRC)}", flush=True)
        try:
            {".pdf": do_pdf, ".pptx": do_pptx, ".xlsx": do_xlsx}.get(ext, lambda p, o: do_video(p, o, not a.no_ocr))(path, out)
        except Exception as ex:
            print(f"   FAILED: {ex}", flush=True)
    print("done", flush=True)


if __name__ == "__main__":
    sys.exit(main())
